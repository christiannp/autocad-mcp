"""Annotation: text, dimensions, leaders and tables."""

from __future__ import annotations

import math
from typing import Any

from .. import com, lisp, util
from ..errors import AcadError
from ..registry import tool

ALIGNMENT = {
    "left": 0, "center": 1, "centre": 1, "right": 2, "aligned": 3, "middle": 4,
    "fit": 5, "top-left": 6, "top-center": 7, "top-centre": 7, "top-right": 8,
    "middle-left": 9, "middle-center": 10, "middle-centre": 10, "middle-right": 11,
    "bottom-left": 12, "bottom-center": 13, "bottom-centre": 13, "bottom-right": 14,
}

ATTACHMENT = {
    "top-left": 1, "top-center": 2, "top-centre": 2, "top-right": 3,
    "middle-left": 4, "middle-center": 5, "middle-centre": 5, "middle-right": 6,
    "bottom-left": 7, "bottom-center": 8, "bottom-centre": 8, "bottom-right": 9,
}


@tool(description="Place single-line text. alignment can be left, center, right, middle, fit, top-left ... bottom-right.")
def draw_text(
    text: str,
    position: list[float],
    height: float = 2.5,
    rotation: float = 0.0,
    alignment: str | None = None,
    style: str | None = None,
    width_factor: float | None = None,
    layer: str | None = None,
    color: Any = None,
    space: str = "auto",
    drawing: str | None = None,
) -> dict[str, Any]:
    if float(height) <= 0:
        raise AcadError("text height must be greater than zero")

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        target = com.space(doc, space)
        if layer:
            util.ensure_layer(doc, layer)
        ent = com.retry(
            lambda: target.AddText(str(text), com.pt(position), float(height))
        )
        if rotation:
            ent.Rotation = math.radians(float(rotation))
        if style:
            ent.StyleName = str(style)
        if width_factor is not None:
            ent.ScaleFactor = float(width_factor)
        if alignment:
            key = str(alignment).strip().lower()
            if key not in ALIGNMENT:
                raise AcadError(
                    f"unknown alignment {alignment!r}; try "
                    "left, center, right, middle, top-left, bottom-right ..."
                )
            ent.Alignment = ALIGNMENT[key]
            if ALIGNMENT[key] != 0:
                ent.TextAlignmentPoint = com.pt(position)
        util.apply_props(ent, layer=layer, colour=color)
        return util.describe(ent)

    return com.run_com(work, timeout=120)


@tool(description=(
    "Place paragraph (multiline) text. Use \\P in the text for a line break. "
    "width is the wrap width; 0 means no wrapping."
))
def draw_mtext(
    text: str,
    position: list[float],
    width: float = 0.0,
    height: float = 2.5,
    rotation: float = 0.0,
    attachment: str | None = None,
    style: str | None = None,
    layer: str | None = None,
    color: Any = None,
    space: str = "auto",
    drawing: str | None = None,
) -> dict[str, Any]:
    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        target = com.space(doc, space)
        if layer:
            util.ensure_layer(doc, layer)
        ent = com.retry(
            lambda: target.AddMText(com.pt(position), float(width), str(text))
        )
        ent.Height = float(height)
        if rotation:
            ent.Rotation = math.radians(float(rotation))
        if style:
            ent.StyleName = str(style)
        if attachment:
            key = str(attachment).strip().lower()
            if key not in ATTACHMENT:
                raise AcadError(f"unknown attachment {attachment!r}")
            ent.AttachmentPoint = ATTACHMENT[key]
            ent.InsertionPoint = com.pt(position)
        util.apply_props(ent, layer=layer, colour=color)
        return util.describe(ent)

    return com.run_com(work, timeout=120)


@tool(description="Change the wording of existing text, mtext, attributes or dimension overrides.")
def text_edit(
    handles: list[str],
    new_text: str,
    drawing: str | None = None,
) -> dict[str, Any]:
    if not handles:
        raise AcadError("no handles given")

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        done, skipped = [], []
        for h in handles:
            ent = com.by_handle(doc, str(h))
            kind = util.dxf_type(ent)
            if kind in ("TEXT", "MTEXT", "ATTRIB", "ATTDEF"):
                ent.TextString = str(new_text)
                done.append(str(h))
            elif kind == "DIMENSION":
                ent.TextOverride = str(new_text)
                done.append(str(h))
            else:
                skipped.append({"handle": str(h), "type": kind})
        out: dict[str, Any] = {"changed": done, "count": len(done)}
        if skipped:
            out["not_text"] = skipped
        return out

    return com.run_com(work, timeout=300)


@tool(description=(
    "Find and replace wording across a drawing's text, mtext, attributes and "
    "dimension overrides. Leave replace_with empty to only report matches."
))
def text_find_replace(
    find: str,
    replace_with: str | None = None,
    match_case: bool = False,
    layer: str | None = None,
    whole_drawing: bool = True,
    drawing: str | None = None,
) -> dict[str, Any]:
    needle = str(find)
    if not needle:
        raise AcadError("nothing to find")

    filters = '(list (cons 0 "TEXT,MTEXT,ATTDEF,DIMENSION,MULTILEADER")'
    if layer:
        filters += f" (cons 8 {lisp.lstr(str(layer))})"
    filters += ")"

    handles = [
        str(h)
        for h in (lisp.evaluate(lisp.raw(f'(acadmcp:handles (ssget "_X" {filters}))'),
                                timeout=300) or [])
    ]

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        hits: list[dict[str, Any]] = []
        changed = 0
        probe = needle if match_case else needle.lower()

        def scan(ent: Any, getter, setter, label: str) -> None:
            nonlocal changed
            current = getter()
            if current is None:
                return
            text = str(current)
            hay = text if match_case else text.lower()
            if probe not in hay:
                return
            entry = {
                "handle": str(ent.Handle),
                "type": util.dxf_type(ent),
                "layer": str(com.quiet(lambda: ent.Layer, "")),
                "text": text[:200],
                "field": label,
            }
            if replace_with is not None:
                if match_case:
                    updated = text.replace(needle, str(replace_with))
                else:
                    out, low, start = [], text.lower(), 0
                    while True:
                        at = low.find(probe, start)
                        if at < 0:
                            out.append(text[start:])
                            break
                        out.append(text[start:at])
                        out.append(str(replace_with))
                        start = at + len(needle)
                    updated = "".join(out)
                setter(updated)
                entry["now"] = updated[:200]
                changed += 1
            hits.append(entry)

        for h in handles:
            try:
                ent = com.by_handle(doc, h)
            except Exception:  # noqa: BLE001
                continue
            kind = util.dxf_type(ent)
            if kind in ("TEXT", "MTEXT", "ATTDEF", "MULTILEADER"):
                scan(ent, lambda e=ent: com.quiet(lambda: e.TextString),
                     lambda v, e=ent: setattr(e, "TextString", v), "text")
            elif kind == "DIMENSION":
                scan(ent, lambda e=ent: com.quiet(lambda: e.TextOverride),
                     lambda v, e=ent: setattr(e, "TextOverride", v), "override")

        # block attributes live on the inserts, not in the ssget above
        att_handles = [
            str(x)
            for x in (
                lisp.evaluate(
                    lisp.raw('(acadmcp:handles (ssget "_X" (list (cons 0 "INSERT"))))'),
                    timeout=300,
                )
                or []
            )
        ]
        for h in att_handles:
            try:
                ref = com.by_handle(doc, h)
            except Exception:  # noqa: BLE001
                continue
            for att in com.unwrap(com.quiet(lambda: ref.GetAttributes(), []) or []):
                scan(att, lambda a=att: com.quiet(lambda: a.TextString),
                     lambda v, a=att: setattr(a, "TextString", v), "attribute")

        return {
            "searched_for": needle,
            "matches": len(hits),
            "replaced": changed if replace_with is not None else None,
            "results": hits[:200],
        }

    result = com.run_com(work, timeout=600)
    return {k: v for k, v in result.items() if v is not None}


@tool(description=(
    "Add a dimension. kind: linear (horizontal/vertical), aligned, angular, "
    "radial, diameter, ordinate or arc. Give the points the kind needs and "
    "text_position for where the dimension line sits."
))
def draw_dimension(
    kind: str = "aligned",
    point1: list[float] | None = None,
    point2: list[float] | None = None,
    text_position: list[float] | None = None,
    vertex: list[float] | None = None,
    center: list[float] | None = None,
    radius_point: list[float] | None = None,
    leader_length: float = 0.0,
    rotation: float = 0.0,
    text_override: str | None = None,
    style: str | None = None,
    layer: str | None = None,
    space: str = "auto",
    drawing: str | None = None,
) -> dict[str, Any]:
    key = str(kind).strip().lower()

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        target = com.space(doc, space)
        if layer:
            util.ensure_layer(doc, layer)

        if key in ("linear", "rotated", "horizontal", "vertical"):
            if not (point1 and point2 and text_position):
                raise AcadError("a linear dimension needs point1, point2 and text_position")
            angle = float(rotation)
            if key == "horizontal":
                angle = 0.0
            elif key == "vertical":
                angle = 90.0
            ent = com.retry(
                lambda: target.AddDimRotated(
                    com.pt(point1), com.pt(point2), com.pt(text_position),
                    math.radians(angle),
                )
            )
        elif key == "aligned":
            if not (point1 and point2 and text_position):
                raise AcadError("an aligned dimension needs point1, point2 and text_position")
            ent = com.retry(
                lambda: target.AddDimAligned(
                    com.pt(point1), com.pt(point2), com.pt(text_position)
                )
            )
        elif key == "angular":
            if not (vertex and point1 and point2 and text_position):
                raise AcadError(
                    "an angular dimension needs vertex, point1, point2 and text_position"
                )
            ent = com.retry(
                lambda: target.AddDimAngular(
                    com.pt(vertex), com.pt(point1), com.pt(point2), com.pt(text_position)
                )
            )
        elif key in ("radial", "radius"):
            if not (center and radius_point):
                raise AcadError("a radial dimension needs center and radius_point")
            ent = com.retry(
                lambda: target.AddDimRadial(
                    com.pt(center), com.pt(radius_point), float(leader_length)
                )
            )
        elif key in ("diameter", "diametric"):
            if not (point1 and point2):
                raise AcadError(
                    "a diameter dimension needs point1 and point2 on opposite sides "
                    "of the circle"
                )
            ent = com.retry(
                lambda: target.AddDimDiametric(
                    com.pt(point1), com.pt(point2), float(leader_length)
                )
            )
        elif key == "ordinate":
            if not (point1 and text_position):
                raise AcadError("an ordinate dimension needs point1 and text_position")
            ent = com.retry(
                lambda: target.AddDimOrdinate(
                    com.pt(point1), com.pt(text_position), str(rotation) == "0"
                )
            )
        elif key == "arc":
            if not (center and point1 and point2 and text_position):
                raise AcadError(
                    "an arc-length dimension needs center, point1, point2 and text_position"
                )
            ent = com.retry(
                lambda: target.AddDimArc(
                    com.pt(center), com.pt(point1), com.pt(point2), com.pt(text_position)
                )
            )
        else:
            raise AcadError(
                "kind must be linear, aligned, angular, radial, diameter, "
                "ordinate or arc"
            )

        if text_override:
            ent.TextOverride = str(text_override)
        if style:
            ent.StyleName = str(style)
        if layer:
            ent.Layer = str(layer)
        return util.describe(ent)

    return com.run_com(work, timeout=180)


@tool(description="Draw a leader with text at the end. points are the leader vertices, from the arrow to the text.")
def draw_leader(
    points: list[list[float]],
    text: str | None = None,
    text_height: float = 2.5,
    arrow: bool = True,
    layer: str | None = None,
    space: str = "auto",
    drawing: str | None = None,
) -> dict[str, Any]:
    if not points or len(points) < 2:
        raise AcadError("a leader needs at least two points")

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        target = com.space(doc, space)
        if layer:
            util.ensure_layer(doc, layer)
        note = None
        if text:
            tip = points[-1]
            note = com.retry(
                lambda: target.AddMText(com.pt(tip), 0.0, str(text))
            )
            note.Height = float(text_height)
            if layer:
                note.Layer = str(layer)
        ent = com.retry(
            lambda: target.AddLeader(com.flat3d(points), note, 0 if arrow else 1)
        )
        if layer:
            ent.Layer = str(layer)
        out = util.describe(ent)
        if note is not None:
            out["text_handle"] = str(note.Handle)
            out["text"] = str(text)
        return out

    return com.run_com(work, timeout=180)


@tool(description=(
    "Insert a table and fill it with data. data is a list of rows, each a list "
    "of cell values; the first row is treated as the header."
))
def draw_table(
    position: list[float],
    data: list[list[Any]],
    row_height: float = 8.0,
    column_width: float = 40.0,
    title: str | None = None,
    layer: str | None = None,
    space: str = "auto",
    drawing: str | None = None,
) -> dict[str, Any]:
    if not data:
        raise AcadError("no data given")
    columns = max(len(r) for r in data)
    header_rows = 1 if title else 0
    rows = len(data) + 1 + header_rows  # AutoCAD reserves a title row

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        target = com.space(doc, space)
        if layer:
            util.ensure_layer(doc, layer)
        table = com.retry(
            lambda: target.AddTable(
                com.pt(position), int(rows), int(columns),
                float(row_height), float(column_width),
            ),
            context="creating the table",
        )
        if title:
            com.quiet(lambda: table.SetText(0, 0, str(title)))
        offset = 1
        for r, row in enumerate(data):
            for c, value in enumerate(row):
                com.quiet(
                    lambda r=r, c=c, v=value: table.SetText(r + offset, c, str(v))
                )
        if layer:
            table.Layer = str(layer)
        return {
            "handle": str(table.Handle),
            "rows": int(rows),
            "columns": int(columns),
            "filled": sum(len(r) for r in data),
        }

    return com.run_com(work, timeout=300)


@tool(description=(
    "Draw a multileader (MLEADER): an arrow from the first point, through any "
    "further points, to a text note. This is the modern leader most drawing "
    "standards expect; draw_leader makes the legacy kind."
))
def draw_mleader(
    points: list[list[float]],
    text: str,
    text_height: float | None = None,
    style: str | None = None,
    landing: bool | None = None,
    layer: str | None = None,
    space: str = "auto",
    drawing: str | None = None,
) -> dict[str, Any]:
    if not points or len(points) < 2:
        raise AcadError("a multileader needs at least two points (arrow, then text)")

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        target = com.space(doc, space)
        if layer:
            util.ensure_layer(doc, layer)
        result = com.retry(lambda: target.AddMLeader(com.flat3d(points), 0))
        ent = result[0] if isinstance(result, (tuple, list)) else result
        ent = com.by_handle(doc, str(ent.Handle))   # re-fetch: late binding
        if style:
            com.quiet(lambda: setattr(ent, "StyleName", str(style)))
        com.quiet(lambda: setattr(ent, "ContentType", 2))   # mtext content
        ent.TextString = str(text)
        if text_height:
            com.quiet(lambda: setattr(ent, "TextHeight", float(text_height)))
        if landing is not None:
            com.quiet(lambda: setattr(ent, "DoglegLength", float(ent.DoglegLength) if landing else 0.0))
        if layer:
            ent.Layer = str(layer)
        out = util.describe(ent)
        out["text"] = str(text)
        return out

    try:
        return com.run_com(work, timeout=180)
    except AcadError as exc:
        # COM's AddMLeader is temperamental; the command always works
        doc = com.run_com(lambda: com.find_doc(drawing), timeout=60) if drawing else None
        steps: list[Any] = []
        if layer:
            steps.append(lisp.raw(f'(setvar "CLAYER" {lisp.lstr(str(layer))})'))
        steps.append(lisp.command("_.MLEADER", *points, str(text)))
        _, created = lisp.capture(
            lisp.pushed({"CLAYER": lisp.raw('(getvar "CLAYER")')}, lisp.progn(*steps)),
            doc=doc, timeout=180,
        )
        return {"created": created, "text": str(text), "via": "MLEADER command", "com_error": str(exc)[:120]}


def _project(p: list[float], origin: list[float], direction: list[float]) -> list[float]:
    """Project p onto the line through origin with unit direction."""
    t = (p[0] - origin[0]) * direction[0] + (p[1] - origin[1]) * direction[1]
    return [origin[0] + direction[0] * t, origin[1] + direction[1] * t]


@tool(description=(
    "A run of dimensions along a line of points (what DIMCONTINUE / DIMBASELINE "
    "produce): one dimension per consecutive pair, all on one dimension line "
    "through dimension_line_point. kind linear measures along `rotation` "
    "degrees (0 horizontal, 90 vertical); aligned follows the points. baseline "
    "stacks every dimension from the first point instead, spacing apart."
))
def dimension_chain(
    points: list[list[float]],
    dimension_line_point: list[float],
    kind: str = "linear",
    rotation: float = 0.0,
    baseline: bool = False,
    spacing: float | None = None,
    style: str | None = None,
    layer: str | None = None,
    space: str = "auto",
    drawing: str | None = None,
) -> dict[str, Any]:
    if not points or len(points) < 2:
        raise AcadError("a chain needs at least two points")
    key = str(kind).strip().lower()
    if key not in ("linear", "aligned", "horizontal", "vertical"):
        raise AcadError("kind must be linear, aligned, horizontal or vertical")
    angle = float(rotation)
    if key == "horizontal":
        angle = 0.0
    elif key == "vertical":
        angle = 90.0
    pts = [[float(p[0]), float(p[1])] for p in points]
    dl = [float(dimension_line_point[0]), float(dimension_line_point[1])]

    if key == "aligned":
        dx, dy = pts[-1][0] - pts[0][0], pts[-1][1] - pts[0][1]
        n = math.hypot(dx, dy) or 1.0
        direction = [dx / n, dy / n]
    else:
        direction = [math.cos(math.radians(angle)), math.sin(math.radians(angle))]
    normal = [-direction[1], direction[0]]

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        target = com.space(doc, space)
        if layer:
            util.ensure_layer(doc, layer)
        gap = float(spacing) if spacing else float(
            com.quiet(lambda: doc.GetVariable("DIMDLI"), 3.75) or 3.75
        ) * float(com.quiet(lambda: doc.GetVariable("DIMSCALE"), 1.0) or 1.0)
        made = []
        pairs = [(pts[0], p) for p in pts[1:]] if baseline else list(zip(pts, pts[1:]))
        for i, (a, b) in enumerate(pairs):
            offset = i * gap if baseline else 0.0
            mid = [(a[0] + b[0]) / 2, (a[1] + b[1]) / 2]
            line_pt = [dl[0] + normal[0] * offset, dl[1] + normal[1] * offset]
            text_pos = _project(mid, line_pt, direction)
            if key == "aligned":
                ent = com.retry(
                    lambda a=a, b=b, t=text_pos: target.AddDimAligned(com.pt(a), com.pt(b), com.pt(t))
                )
            else:
                ent = com.retry(
                    lambda a=a, b=b, t=text_pos: target.AddDimRotated(
                        com.pt(a), com.pt(b), com.pt(t), math.radians(angle)
                    )
                )
            if style:
                com.quiet(lambda e=ent: setattr(e, "StyleName", str(style)))
            if layer:
                ent.Layer = str(layer)
            made.append({
                "handle": str(ent.Handle),
                "measurement": com.quiet(lambda e=ent: round(float(e.Measurement), 6)),
            })
        return {"created": [m["handle"] for m in made], "dimensions": made, "count": len(made),
                "baseline": bool(baseline)}

    return com.run_com(work, timeout=300)


@tool(description=(
    "Change existing dimensions: override the text (use '' to go back to the "
    "measured value, or '<>' inside text to keep the number, e.g. '<> TYP.'), "
    "move the text, change the style, overall scale, text height, arrowhead "
    "size or decimal places. Omit everything to just read them."
))
def dimension_edit(
    handles: list[str],
    text_override: str | None = None,
    text_position: list[float] | None = None,
    style: str | None = None,
    scale: float | None = None,
    text_height: float | None = None,
    arrow_size: float | None = None,
    precision: int | None = None,
    text_rotation: float | None = None,
    drawing: str | None = None,
) -> dict[str, Any]:
    if not handles:
        raise AcadError("no handles given")

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        rows = []
        for h in handles:
            ent = com.by_handle(doc, str(h))
            if util.dxf_type(ent) != "DIMENSION":
                rows.append({"handle": str(h), "skipped": f"{util.dxf_type(ent)} is not a dimension"})
                continue
            if text_override is not None:
                ent.TextOverride = str(text_override)
            if text_position is not None:
                ent.TextPosition = com.pt(text_position)
            if style:
                ent.StyleName = str(style)
            if scale is not None:
                ent.ScaleFactor = float(scale)
            if text_height is not None:
                ent.TextHeight = float(text_height)
            if arrow_size is not None:
                ent.ArrowheadSize = float(arrow_size)
            if precision is not None:
                ent.PrimaryUnitsPrecision = int(precision)
            if text_rotation is not None:
                ent.TextRotation = math.radians(float(text_rotation))
            q = com.quiet
            rows.append({
                "handle": str(ent.Handle),
                "measurement": q(lambda: round(float(ent.Measurement), 6)),
                "text_override": q(lambda: str(ent.TextOverride)),
                "text_position": q(lambda: util.round_pt(ent.TextPosition)),
                "style": q(lambda: str(ent.StyleName)),
                "scale": q(lambda: round(float(ent.ScaleFactor), 6)),
                "text_height": q(lambda: round(float(ent.TextHeight), 6)),
            })
        return {"count": len(rows), "dimensions": rows}

    return com.run_com(work, timeout=300)


@tool(readonly=True, description=(
    "Read an existing table: every cell's text as rows, plus the size of the "
    "grid, so a schedule in the drawing can be checked or copied."
))
def table_read(handle: str, drawing: str | None = None) -> dict[str, Any]:
    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        table = com.by_handle(doc, str(handle))
        if util.dxf_type(table) != "ACAD_TABLE":
            raise AcadError(f"{handle} is a {util.dxf_type(table)}, not a table")
        rows, cols = int(table.Rows), int(table.Columns)
        grid = [
            [str(com.quiet(lambda r=r, c=c: table.GetText(r, c), "") or "") for c in range(cols)]
            for r in range(rows)
        ]
        return {
            "handle": str(table.Handle),
            "rows": rows,
            "columns": cols,
            "style": com.quiet(lambda: str(table.StyleName)),
            "position": com.quiet(lambda: util.round_pt(table.InsertionPoint)),
            "column_widths": [com.quiet(lambda c=c: round(float(table.GetColumnWidth(c)), 4)) for c in range(cols)],
            "row_heights": [com.quiet(lambda r=r: round(float(table.GetRowHeight(r)), 4)) for r in range(rows)],
            "cells": grid,
        }

    return com.run_com(work, timeout=180)


@tool(description=(
    "Edit an existing table: set cells ([{row, col, value}], 0-based), insert "
    "or delete rows/columns, resize columns or rows, merge a block of cells, "
    "change the table style. update_data_links refreshes every Excel-linked "
    "table in the drawing (DATALINKUPDATE) - use it with no handle."
))
def table_edit(
    handle: str | None = None,
    cells: list[dict[str, Any]] | None = None,
    insert_rows: dict[str, Any] | None = None,
    delete_rows: list[int] | None = None,
    insert_columns: dict[str, Any] | None = None,
    delete_columns: list[int] | None = None,
    column_widths: dict[str, float] | None = None,
    row_heights: dict[str, float] | None = None,
    merge: dict[str, int] | None = None,
    style: str | None = None,
    update_data_links: bool = False,
    drawing: str | None = None,
) -> dict[str, Any]:
    if update_data_links:
        doc = com.run_com(lambda: com.find_doc(drawing), timeout=60) if drawing else None
        lisp.run_command("_.DATALINKUPDATE", "_Update", "_K", doc=doc, timeout=300)
        if not handle:
            return {"data_links": "updated"}
    if not handle:
        raise AcadError("give the table's handle")

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        table = com.by_handle(doc, str(handle))
        if util.dxf_type(table) != "ACAD_TABLE":
            raise AcadError(f"{handle} is a {util.dxf_type(table)}, not a table")
        done: dict[str, Any] = {}
        if style:
            table.StyleName = str(style)
            done["style"] = str(style)
        if insert_rows:
            at = int(insert_rows.get("index", int(table.Rows)))
            n = int(insert_rows.get("count", 1))
            h = float(insert_rows.get("height") or table.GetRowHeight(max(0, min(at, int(table.Rows) - 1))))
            com.retry(lambda: table.InsertRows(at, h, n))
            done["inserted_rows"] = n
        if insert_columns:
            at = int(insert_columns.get("index", int(table.Columns)))
            n = int(insert_columns.get("count", 1))
            w = float(insert_columns.get("width") or table.GetColumnWidth(max(0, min(at, int(table.Columns) - 1))))
            com.retry(lambda: table.InsertColumns(at, w, n))
            done["inserted_columns"] = n
        for row in cells or []:
            table.SetText(int(row["row"]), int(row["col"]), str(row.get("value", "")))
        if cells:
            done["cells_set"] = len(cells)
        for r in sorted({int(r) for r in (delete_rows or [])}, reverse=True):
            com.retry(lambda r=r: table.DeleteRows(r, 1))
        if delete_rows:
            done["deleted_rows"] = len(delete_rows)
        for c in sorted({int(c) for c in (delete_columns or [])}, reverse=True):
            com.retry(lambda c=c: table.DeleteColumns(c, 1))
        if delete_columns:
            done["deleted_columns"] = len(delete_columns)
        for c, w in (column_widths or {}).items():
            table.SetColumnWidth(int(c), float(w))
        for r, hgt in (row_heights or {}).items():
            table.SetRowHeight(int(r), float(hgt))
        if merge:
            com.retry(lambda: table.MergeCells(
                int(merge["top"]), int(merge["bottom"]), int(merge["left"]), int(merge["right"])
            ))
            done["merged"] = merge
        com.quiet(lambda: table.RecomputeTableBlock(True))
        done.update({"handle": str(table.Handle), "rows": int(table.Rows), "columns": int(table.Columns)})
        return done

    return com.run_com(work, timeout=300)


@tool(description=(
    "Insert a table filled from a spreadsheet: an .xlsx sheet (optionally a "
    "cell range like A1:D20) or a .csv. The values are copied in, not linked - "
    "AutoCAD's live Excel DATALINK can only be created in its own dialog; if "
    "she needs one, create it there once and table_edit(update_data_links) "
    "refreshes it."
))
def table_from_spreadsheet(
    path: str,
    position: list[float],
    sheet: str | None = None,
    cell_range: str | None = None,
    header_rows: int = 1,
    title: str | None = None,
    row_height: float = 8.0,
    column_width: float = 40.0,
    layer: str | None = None,
    space: str = "auto",
    drawing: str | None = None,
) -> dict[str, Any]:
    import os

    source = os.path.abspath(os.path.expanduser(str(path)))
    if not os.path.isfile(source):
        raise AcadError(f"no such file: {source}")
    data: list[list[Any]] = []
    if source.lower().endswith(".csv"):
        import csv

        with open(source, newline="", encoding="utf-8-sig") as fh:
            data = [list(r) for r in csv.reader(fh)]
    else:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise AcadError("openpyxl is not installed, so .xlsx files cannot be read") from exc
        wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb.active
        rows_iter = ws[cell_range] if cell_range else ws.iter_rows()
        for r in rows_iter:
            data.append(["" if c.value is None else c.value for c in r])
        wb.close()
    data = [r for r in data if any(str(v).strip() for v in r)]
    if not data:
        raise AcadError("the spreadsheet range is empty")
    out = draw_table(
        position, data, row_height=row_height, column_width=column_width,
        title=title, layer=layer, space=space, drawing=drawing,
    )
    out["source"] = source
    out["header_rows"] = int(header_rows)
    return out


@tool(description=(
    "Combine several single-line TEXT objects into one MTEXT paragraph "
    "(TXT2MTXT), keeping their order top to bottom."
))
def text_combine(handles: list[str], drawing: str | None = None) -> dict[str, Any]:
    if len(handles) < 1:
        raise AcadError("no handles given")
    doc = com.run_com(lambda: com.find_doc(drawing), timeout=60) if drawing else None
    _, created = lisp.capture(
        lisp.command("_.TXT2MTXT", lisp.ss_from([str(h) for h in handles]), ""),
        doc=doc, timeout=180,
    )
    return {"combined": len(handles), "created": created}


@tool(description=(
    "Annotation scales. list: the drawing's scale list and the current scale "
    "(CANNOSCALE). set_current: make a scale current for new annotative "
    "objects. add: add a scale to the list (paper_units:drawing_units, e.g. "
    "1:100). add_to_objects / remove_from_objects: give annotative objects "
    "(handles) a scale representation."
))
def annotation_scale(
    action: str = "list",
    scale: str | None = None,
    paper_units: float = 1.0,
    drawing_units: float | None = None,
    handles: list[str] | None = None,
    drawing: str | None = None,
) -> dict[str, Any]:
    verb = str(action).strip().lower()
    doc = com.run_com(lambda: com.find_doc(drawing), timeout=60) if drawing else None

    def scales() -> dict[str, Any]:
        rows = lisp.evaluate(
            lisp.raw(
                '(progn (setq amx-out nil) (foreach amx-p (dictsearch (namedobjdict) "ACAD_SCALELIST") '
                "(if (= (car amx-p) 350) (progn (setq amx-e (entget (cdr amx-p))) "
                "(setq amx-out (cons (list (cdr (assoc 300 amx-e)) (cdr (assoc 140 amx-e)) "
                "(cdr (assoc 141 amx-e))) amx-out))))) (reverse amx-out))"
            ),
            doc=doc, timeout=60,
        ) or []
        return {
            "current": lisp.evaluate(lisp.raw('(getvar "CANNOSCALE")'), doc=doc, timeout=60),
            "current_value": lisp.evaluate(lisp.raw('(getvar "CANNOSCALEVALUE")'), doc=doc, timeout=60),
            "scales": [{"name": r[0], "paper": r[1], "drawing": r[2]} for r in rows if isinstance(r, list)],
        }

    if verb == "list":
        return scales()
    if not scale:
        raise AcadError(f"{verb} needs the scale name, e.g. '1:100'")
    if verb == "set_current":
        lisp.evaluate(lisp.raw(f'(setvar "CANNOSCALE" {lisp.lstr(str(scale))})'), doc=doc, timeout=60)
        return {"current": scale}
    if verb == "add":
        ratio = f"{paper_units:g}:{drawing_units:g}" if drawing_units else str(scale)
        lisp.run_command("_.-SCALELISTEDIT", "_Add", str(scale), ratio, "_Exit", doc=doc, timeout=60)
        return {"added": scale, "ratio": ratio}
    if verb in ("add_to_objects", "remove_from_objects"):
        if not handles:
            raise AcadError(f"{verb} needs handles")
        option = "_Add" if verb == "add_to_objects" else "_Delete"
        lisp.run_command(
            "_.-OBJECTSCALE", lisp.ss_from([str(h) for h in handles]), "", option, str(scale), "",
            doc=doc, timeout=120,
        )
        return {verb: scale, "count": len(handles)}
    raise AcadError("action must be list, set_current, add, add_to_objects or remove_from_objects")

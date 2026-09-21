"""Understanding a drawing: statistics and data extraction."""

from __future__ import annotations

import os
from collections import Counter, OrderedDict
from datetime import datetime
from typing import Any

from .. import com, lisp, util
from ..errors import AcadError
from ..registry import tool


@tool(readonly=True, description=(
    "A full picture of a drawing in one call: entity counts by type and layer, "
    "layers, blocks, layouts, xrefs, text and dimension styles, units and extents. "
    "Use this first when meeting an unfamiliar drawing."
))
def drawing_info(
    include_layers: bool = True,
    include_blocks: bool = True,
    drawing: str | None = None,
) -> dict[str, Any]:
    rows = lisp.evaluate(
        lisp.raw(
            '(progn (setq amq-ss (ssget "_X") amq-i 0 amq-out nil)'
            " (if amq-ss (repeat (sslength amq-ss) (setq amq-e (entget (ssname amq-ss amq-i))"
            " amq-out (cons (list (cdr (assoc 0 amq-e)) (cdr (assoc 8 amq-e))"
            " (cdr (assoc 410 amq-e))) amq-out) amq-i (1+ amq-i))))"
            " (reverse amq-out))"
        ),
        timeout=600,
    ) or []

    by_type: Counter = Counter()
    by_layer: Counter = Counter()
    by_space: Counter = Counter()
    for row in rows:
        if not isinstance(row, list) or len(row) < 3:
            continue
        by_type[str(row[0])] += 1
        by_layer[str(row[1])] += 1
        by_space[str(row[2] or "Model")] += 1

    def work() -> dict[str, Any]:
        doc = com.find_doc(drawing)
        units = int(com.quiet(lambda: doc.GetVariable("INSUNITS"), 0) or 0)
        from .session import UNIT_NAMES

        info: dict[str, Any] = {
            "name": str(doc.Name),
            "path": str(com.quiet(lambda: doc.FullName, "") or "(never saved)"),
            "saved": bool(com.quiet(lambda: doc.Saved, True)),
            "units": UNIT_NAMES.get(units, units),
            "total_entities": sum(by_type.values()),
            "entities_by_type": dict(by_type.most_common()),
            "entities_by_space": dict(by_space.most_common()),
            "layer_count": int(com.quiet(lambda: doc.Layers.Count, 0) or 0),
            "current_layer": str(com.quiet(lambda: doc.ActiveLayer.Name, "")),
            "layouts": [
                str(doc.Layouts.Item(i).Name) for i in range(int(doc.Layouts.Count))
            ],
            "text_styles": int(com.quiet(lambda: doc.TextStyles.Count, 0) or 0),
            "dim_styles": int(com.quiet(lambda: doc.DimStyles.Count, 0) or 0),
            "current_dim_style": str(com.quiet(lambda: doc.ActiveDimStyle.Name, "")),
        }
        if by_layer:
            info["entities_by_layer"] = dict(by_layer.most_common(40))
        if include_layers:
            info["layers"] = [
                {
                    "name": str(doc.Layers.Item(i).Name),
                    "color": com.quiet(lambda i=i: int(doc.Layers.Item(i).Color)),
                    "on": bool(com.quiet(lambda i=i: doc.Layers.Item(i).LayerOn, True)),
                    "frozen": bool(com.quiet(lambda i=i: doc.Layers.Item(i).Freeze, False)),
                    "locked": bool(com.quiet(lambda i=i: doc.Layers.Item(i).Lock, False)),
                    "objects": by_layer.get(str(doc.Layers.Item(i).Name), 0),
                }
                for i in range(int(doc.Layers.Count))
            ]
        if include_blocks:
            blocks, xrefs = [], []
            for i in range(int(doc.Blocks.Count)):
                block = doc.Blocks.Item(i)
                name = str(block.Name)
                if name.startswith("*") or com.quiet(lambda: block.IsLayout, False):
                    continue
                if com.quiet(lambda: block.IsXRef, False):
                    path = str(com.quiet(lambda: block.Path, "") or "")
                    xrefs.append({"name": name, "path": path, "found": os.path.isfile(path)})
                else:
                    blocks.append(name)
            info["blocks"] = sorted(blocks)
            info["xrefs"] = xrefs

        lo = com.quiet(lambda: com.unwrap(doc.GetVariable("EXTMIN")))
        hi = com.quiet(lambda: com.unwrap(doc.GetVariable("EXTMAX")))
        if lo and hi and all(abs(v) < 1e19 for v in list(lo) + list(hi)):
            info["extents"] = {
                "min": [round(float(v), 4) for v in lo],
                "max": [round(float(v), 4) for v in hi],
                "size": [round(float(hi[i]) - float(lo[i]), 4) for i in range(3)],
            }
        return info

    return com.run_com(work, timeout=300)


@tool(description=(
    "Extract drawing data to a spreadsheet or CSV - the scriptable replacement "
    "for AutoCAD's Data Extraction wizard. Pulls block attributes and geometry "
    "for whatever the filters match, one row per object."
))
def data_extract(
    path: str,
    block: str | None = None,
    type: str | None = None,
    layer: str | None = None,
    space: str | None = None,
    include_geometry: bool = True,
    group_identical: bool = False,
    drawing: str | None = None,
) -> dict[str, Any]:
    target = os.path.abspath(os.path.expanduser(str(path)))
    os.makedirs(os.path.dirname(target) or ".", exist_ok=True)
    ext = os.path.splitext(target)[1].lower()
    if ext not in (".xlsx", ".csv"):
        target += ".xlsx"
        ext = ".xlsx"

    from .select import _filter_lisp, _filter_pairs

    pairs = _filter_pairs(
        type or ("INSERT" if block else None), layer, None, block, None, None,
        space, None, None,
    )
    handles = [
        str(h)
        for h in (
            lisp.evaluate(
                lisp.raw(f'(acadmcp:handles (ssget "_X" {_filter_lisp(pairs)}))'),
                timeout=600,
            )
            or []
        )
    ]
    if not handles:
        raise AcadError("Nothing matched those filters, so there is nothing to extract.")

    def work() -> list[dict[str, Any]]:
        doc = com.find_doc(drawing)
        records: list[dict[str, Any]] = []
        for h in handles:
            try:
                ent = com.by_handle(doc, h)
            except Exception:  # noqa: BLE001
                continue
            kind = util.dxf_type(ent)
            row: OrderedDict[str, Any] = OrderedDict()
            row["Handle"] = h
            row["Type"] = kind
            row["Layer"] = str(com.quiet(lambda: ent.Layer, ""))
            if kind == "INSERT":
                row["Block"] = str(com.quiet(lambda: ent.Name, ""))
            for att in com.unwrap(com.quiet(lambda: ent.GetAttributes(), []) or []):
                row[str(att.TagString)] = str(att.TextString)
            if kind in ("TEXT", "MTEXT"):
                row["Text"] = str(com.quiet(lambda: ent.TextString, ""))
            if include_geometry:
                point = (
                    com.quiet(lambda: util.round_pt(ent.InsertionPoint))
                    or com.quiet(lambda: util.round_pt(ent.Center))
                    or com.quiet(lambda: util.round_pt(ent.StartPoint))
                )
                if point:
                    row["X"], row["Y"] = point[0], point[1]
                    if len(point) > 2:
                        row["Z"] = point[2]
                area = com.quiet(lambda: round(float(ent.Area), 4))
                if area is not None:
                    row["Area"] = area
                length = com.quiet(lambda: round(float(ent.Length), 4))
                if length is not None:
                    row["Length"] = length
                rot = com.quiet(lambda: round(float(ent.Rotation) * 180 / 3.141592653589793, 4))
                if rot is not None:
                    row["Rotation"] = rot
            records.append(dict(row))
        return records

    records = com.run_com(work, timeout=900)

    columns: list[str] = []
    for record in records:
        for key in record:
            if key not in columns:
                columns.append(key)

    if group_identical:
        ignore = {"Handle", "X", "Y", "Z"}
        counter: Counter = Counter()
        keys = [c for c in columns if c not in ignore]
        for record in records:
            counter[tuple(record.get(k, "") for k in keys)] += 1
        columns = keys + ["Count"]
        records = [
            dict(zip(keys, values)) | {"Count": count}
            for values, count in counter.most_common()
        ]

    if ext == ".csv":
        import csv

        with open(target, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(records)
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        sheet.title = "Extract"
        sheet.append(columns)
        head_fill = PatternFill("solid", fgColor="1F3864")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        for record in records:
            sheet.append([record.get(c, "") for c in columns])
        for index, name in enumerate(columns, start=1):
            widest = max([len(str(name))] + [len(str(r.get(name, ""))) for r in records[:400]])
            sheet.column_dimensions[get_column_letter(index)].width = min(42, widest + 3)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        meta = book.create_sheet("About")
        meta.append(["Extracted from", str(records and "drawing" or "")])
        meta.append(["When", datetime.now().strftime("%Y-%m-%d %H:%M")])
        meta.append(["Filters", f"block={block} type={type} layer={layer} space={space}"])
        meta.append(["Rows", len(records)])
        meta.column_dimensions["A"].width = 18
        meta.column_dimensions["B"].width = 60

        book.save(target)

    return {
        "file": target,
        "rows": len(records),
        "columns": columns,
        "grouped": bool(group_identical),
        "size_bytes": os.path.getsize(target) if os.path.isfile(target) else 0,
    }
